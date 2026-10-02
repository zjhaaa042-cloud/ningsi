"""按 docs/DOC_ALIGNMENT.md 修订两份交付 docx：补封面、统一口径、补表格与篇幅。

实现方式：python-docx + lxml，直接复制既有段落的 pPr/rPr 以保证版式一致；
新增标题沿用原文档标题段落的格式，因此仍会出现在目录中（Word 打开时自动更新域）。
"""

from __future__ import annotations

import shutil
import time
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[2]          # D:/Desktop/ninsi
SOURCE = ROOT
OUT = ROOT / "deliverables"
OUT.mkdir(parents=True, exist_ok=True)

LOG = []


def log(message: str) -> None:
    LOG.append(message)
    print(message)


# ---------- 基础工具 ----------
def iter_paragraphs(doc):
    for paragraph in doc.paragraphs:
        yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def set_paragraph_text(paragraph, text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(text)
        return
    paragraph.runs[0].text = text
    for run in paragraph.runs[1:]:
        run.text = ""


def find(doc, *, exact=None, contains=None, startswith=None, after=None):
    seen = after is None
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if not seen:
            if text == after:
                seen = True
            continue
        if exact is not None and text == exact:
            return paragraph
        if contains is not None and contains in text:
            return paragraph
        if startswith is not None and text.startswith(startswith):
            return paragraph
    return None


def replace_all(doc, pairs) -> None:
    for old, new in pairs:
        hits = 0
        for paragraph in iter_paragraphs(doc):
            text = paragraph.text
            if old in text:
                set_paragraph_text(paragraph, text.replace(old, new))
                hits += 1
        log(("  [替换] " if hits else "  [缺失] ") + repr(old[:26]) + f" -> {hits} 处")


def clone_paragraph(template, text: str):
    element = deepcopy(template._p)
    for child in list(element):
        if child.tag != qn("w:pPr"):
            element.remove(child)
    rpr = None
    if template.runs:
        rpr = template.runs[0]._r.find(qn("w:rPr"))
    run = OxmlElement("w:r")
    if rpr is not None:
        run.append(deepcopy(rpr))
    node = OxmlElement("w:t")
    node.set(qn("xml:space"), "preserve")
    node.text = text
    run.append(node)
    element.append(run)
    return element


def build_table(doc, rows, font_size=9):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.text = str(value)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(font_size)
    return table._tbl


def insert_block(anchor_paragraph, elements, before=False) -> None:
    """把若干 XML 元素插到锚点段落之前或之后，保持给定顺序。"""
    cursor = anchor_paragraph._p
    for element in elements:
        if before:
            cursor.addprevious(element)
        else:
            cursor.addnext(element)
            cursor = element


def spacer(doc):
    return clone_paragraph(doc.paragraphs[0], "")


COVER_PREFIXES = ("【A09】", "命题企业：", "产品名称：", "产品 slogan：", "文档类型：",
                  "版本：v2", "团队名称：", "日期：2026")

PPR_SUCCESSORS = ("w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap",
                  "w:jc", "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
                  "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")


def prepare_cover(doc, doc_type: str) -> None:
    """文档自带整页封面图：只做几何修正，不再另插文字封面页。

    修正点：删除此前插入的文字封面；把整页图高度收到 29.5 cm 并等比缩放（避免超出页面被裁剪或挤出空白页）；
    段落前后间距归零、关闭网格吸附、段落标记字号设为 1pt，使封面严格占满第一页且顶部不被压边。
    """
    cover = None
    for paragraph in doc.paragraphs[:8]:
        for extent in paragraph._p.findall(".//" + qn("wp:extent")):
            if int(extent.get("cy", "0")) >= 10_000_000:
                cover = (paragraph, extent)
                break
        if cover is not None:
            break
    if cover is None:
        log("  [封面] 未找到整页封面图，跳过")
        return

    paragraph, extent = cover
    removed = 0
    for candidate in list(doc.paragraphs):
        if candidate._p is paragraph._p:
            break
        if candidate.text.strip().startswith(COVER_PREFIXES):
            candidate._p.getparent().remove(candidate._p)
            removed += 1

    cx, cy = int(extent.get("cx")), int(extent.get("cy"))
    target_cy = 10620000                     # 29.5 cm，留 2 mm 安全余量
    if cy > target_cy:
        new_cx = int(round(cx * target_cy / cy))
        for node in [extent] + paragraph._p.findall(".//" + qn("a:ext")):
            node.set("cx", str(new_cx))
            node.set("cy", str(target_cy))
    else:
        new_cx = cx

    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pPr = paragraph._p.get_or_add_pPr()
    existing = pPr.find(qn("w:snapToGrid"))
    if existing is not None:
        pPr.remove(existing)
    snap = OxmlElement("w:snapToGrid")
    snap.set(qn("w:val"), "0")
    pPr.insert_element_before(snap, *PPR_SUCCESSORS)
    mark = pPr.find(qn("w:rPr"))
    if mark is None:
        mark = OxmlElement("w:rPr")
        pPr.insert_element_before(mark, "w:sectPr", "w:pPrChange")
    for tag in ("w:sz", "w:szCs"):
        node = OxmlElement(tag)
        node.set(qn("w:val"), "2")
        mark.append(node)
    log(f"  [封面] 保留文档自带整页封面（{doc_type}）：清理旧文字封面 {removed} 段，"
        f"图幅 {cx / 360000:.2f}×{cy / 360000:.2f}cm → {new_cx / 360000:.2f}×{target_cy / 360000:.2f}cm，"
        f"并关闭网格吸附、段落标记 1pt")


def set_update_fields(doc) -> None:
    settings = doc.settings.element
    element = settings.find(qn("w:updateFields"))
    if element is None:
        element = OxmlElement("w:updateFields")
        settings.append(element)
    element.set(qn("w:val"), "true")
    log("  [目录] 已设置打开时自动更新域（TOC 会刷新为新页码）")


SUFFIX = ""          # 由 main() 按命令行参数设置，例如 "-更新版"


def save(doc, name: str) -> Path:
    if SUFFIX:
        stem, dot, extension = name.rpartition(".")
        name = f"{stem}{SUFFIX}{dot}{extension}"
        target = OUT / name
        doc.save(target)
        log(f"  [保存] {target}")
        return target
    target = OUT / name
    for attempt in range(20):                      # 文件可能正在 Word 中打开，等待其释放
        try:
            doc.save(target)
            log(f"  [保存] {target}")
            return target
        except PermissionError:
            if attempt == 0:
                log(f"  [等待] {target.name} 正被占用，等待释放（最多 100 秒）…")
            time.sleep(5)
    target = target.with_name(f"{target.stem}-更新版{target.suffix}")
    doc.save(target)
    log(f"  [保存] 原文件始终被占用，改写到 {target.name}")
    return target


# ---------- 概要介绍 ----------
OVERVIEW_REPLACEMENTS = [
    (
        "在生活与工作节奏不断加快的背景下，焦虑、情绪低落与注意力难以集中已成为相当普遍的状态。",
        "在生活与工作节奏加快的背景下，焦虑、情绪低落与注意力难以集中已相当普遍。",
    ),
    ("本项目构建凝思，一个", "本项目构建凝思（Ningsi）——「每一段专注，都看得见」——一个"),
    (
        "全部采用可离线运行的开源技术栈",
        "全部采用可离线运行的技术栈（自研代码 + NumPy 等开源依赖；上游采集工程与设备厂商的许可另行确认）",
    ),
    (
        "静息基线模块（闭眼 3 分钟、睁眼 2 分钟）",
        "静息基线模块（上游 M0 协议：闭眼 3 分钟、睁眼 2 分钟）",
    ),
]

OVERVIEW_FOREWORD_CONDENSED = (
    "团队已有的工程正落在这三件事的工程侧：BSense 采集工程套件实现了设备质量检查、多路信号同步采集、"
    "标准容器落盘、在线质量门控与逐窗口质量报告，并提供认知负荷任务（0/1/2-back，每等级 3 区块 × 60 刺激、"
    "目标刺激占 25%）、静息基线协议与六类动作伪迹验证；按被试划分的数据集构建与校验流程也已沉淀。"
    "这些能力原先服务于岗位前状态筛查，尚未形成面向大众的专注力训练闭环；本赛题要求的方向，"
    "正是把它们从「采集与评估」延伸到「评估与强化」。"
)

OVERVIEW_TABLE1 = [
    ["层级", "主要职责", "对应实现模块", "用户可见入口"],
    ["用户层", "学生与职场人群、个人用户、心理教师与辅导员、员工关怀与研究人员四类角色", "报告与趋势按角色区分可见范围", "会话设置、报告查看"],
    ["交互层", "设备接入与质检、评估与量表、实时监测、训练主界面、报告与趋势", "app/ui.py 六个页签；CLI 提供同一套能力的脚本入口", "状态指标、实时波形与频谱、热力图、训练、报告、趋势"],
    ["智能层", "四级处理链 → Welch 频谱 → 频带特征 → 三指标 → 质量门控 → 联合评估 → 训练建议", "signal/、behavior/、assessment/、models/、training/", "指标面板、热力图、评估报告、训练反馈"],
    ["支撑层", "采集抽象、口径与版本管理、可复算证据回填、历史与趋势", "acquisition/、config.py、assessment/report.py、monitoring/history.py", "报告中的证据回填、趋势页"],
    ["安全保障（贯通各层）", "匿名编号、受限目录、不作医疗诊断、结论只给操作建议", "参与者编号约定、报告边界声明与配置中的提示文本", "报告边界说明与预警提示"],
]

OVERVIEW_TABLE2 = [
    ["分类", "工具与组件", "版本与用途"],
    ["运行环境", "Python、NumPy", "Python 3.11–3.13 + NumPy >=1.26；全部计算本地完成，训练现场不依赖公网"],
    ["采集接入", "pylsl（可选）", "1.18.2；通过 LSL 接入任意厂商原始 EEG 流，未接设备时使用内置仿真源联调"],
    ["信号处理", "自研零依赖实现", "Welch 4 秒窗 / 2 秒分段 / 50% 重叠 / 汉宁窗、NFFT 512、0.5–45 Hz 带通与 50/60 Hz 陷波，不依赖 SciPy"],
    ["特征与指标", "signal/indicators.py", "β/θ、α/β、θ/α 相对个体基线做 z 标准化后映射为 0–1 评分（indicator-v1）"],
    ["模型", "自研逻辑回归", "6 维频带特征、被试级 6:2:2 划分、AUC 与混淆矩阵、JSON 版本化保存（logistic-v1）"],
    ["界面", "标准库 Tkinter", "六个页签桌面界面，无需额外依赖即可分发"],
    ["报告与图表", "自研 Markdown / JSON / SVG", "评估报告、状态热力图与趋势图可直接嵌入答辩材料"],
]


def revise_overview(source: Path) -> Path:
    log(f"\n=== 概要介绍：{source.name} ===")
    doc = Document(str(source))
    prepare_cover(doc, "项目概要介绍")
    replace_all(doc, OVERVIEW_REPLACEMENTS)

    condensed = find(doc, contains="团队此前已完成的工作正落在这三件事的工程侧")
    if condensed is not None:
        set_paragraph_text(condensed, OVERVIEW_FOREWORD_CONDENSED)
        log("  [前言] 已把团队已有工程段落压缩为一段")
    else:
        log("  [缺失] 未找到团队已有工程段落")

    trailing = find(doc, contains="本赛题要求的方向，正是把这些能力从")
    if trailing is not None:
        element = trailing._p
        element.getparent().remove(element)
        log("  [前言] 已删除与上一段重复的收尾句（篇幅配比：前言压缩）")

    body_template = find(doc, contains="系统自下而上分为四层")
    caption_template = find(doc, exact="图 2 用户使用流程图")
    anchor = caption_template or body_template
    if anchor is not None and body_template is not None and caption_template is not None:
        block = [
            clone_paragraph(body_template, "七步中的第 ② 步（设备质检）与第 ④ 步（静息基线）是可信度的两个锚点："
                                          "质检决定这次采集能不能用，基线决定指标与谁比。系统为每一步保留可复核记录——"
                                          "质检结果、量表作答、基线统计量、逐窗质量与排除原因、指标与参数版本随会话落盘，"
                                          "结论因此可以被追溯，而不是只能被相信。"),
            clone_paragraph(body_template, "四个层级与用户实际能看到的入口一一对应，如表 1 所示。"),
            clone_paragraph(caption_template, "表 1 四层架构与实现模块对应表"),
            build_table(doc, OVERVIEW_TABLE1),
            spacer(doc),
        ]
        insert_block(anchor, block)
        log("  [新增] 三、功能简介：1 张架构对应表 + 2 段说明")
    else:
        log("  [缺失] 未找到功能简介插入锚点")

    tail_anchor = find(doc, contains="训练现场不依赖公网")
    if tail_anchor is not None and body_template is not None and caption_template is not None:
        block = [
            clone_paragraph(body_template, "工具链与指标口径是同一份定义的两种呈现：代码中的 config.py 集中声明频谱、"
                                          "指标、基线、训练与预警参数，报告再把对应版本号写进结论，"
                                          "因此文档里写的与系统里算的始终是同一套口径。"),
            clone_paragraph(body_template, "各层工具与口径版本的对应关系如表 2 所示。"),
            clone_paragraph(caption_template, "表 2 开发工具与口径版本对照表"),
            build_table(doc, OVERVIEW_TABLE2),
            spacer(doc),
        ]
        insert_block(tail_anchor, block)
        log("  [新增] 五、开发工具与技术：1 张工具口径表 + 2 段说明")
    else:
        log("  [缺失] 未找到开发工具章节锚点")

    set_update_fields(doc)
    return save(doc, "A09-凝思-项目概要介绍-v2.docx")


# ---------- 详细方案 ----------
DETAIL_REPLACEMENTS = [
    ("FocusFlow", "Ningsi"),
    ("带通滤波 1–45 Hz", "带通滤波 0.5–45 Hz"),
    ("表 6.9 三类目标标签与数据量口径（待开发）", "表 6.9 三类目标标签与数据量口径（压力与情绪标签待采集）"),
    ("表 6.11 状态热力图颜色分档与含义（待开发）", "表 6.11 状态热力图颜色分档与含义"),
    ("模型处于待开发状态时", "模型概率尚未接入色阶时（当前版本）"),
    (
        "自然呼吸稳定 30 秒、闭眼 3 分钟、睁眼 2 分钟",
        "自然呼吸稳定 30 秒、闭眼 3 分钟、睁眼 2 分钟（上游 M0 协议；本产品会话内基线为睁眼、闭眼各 2 分钟）",
    ),
    (
        "睁眼静息与闭眼静息两段（默认各 2 分钟，可配置为 1–3 分钟）",
        "睁眼静息与闭眼静息两段（默认各 2 分钟，可配置为 1–3 分钟；两段各自独立质检与建立基线，任务态指标以睁眼基线为参照）",
    ),
    (
        "1–45 Hz 零相位带通滤波",
        "0.5–45 Hz 零相位带通滤波（0.5 Hz 高通 + 45 Hz 低通）",
    ),
    (
        "支持 OpenBCI、Emotiv、NeuroSky 及任何提供 LSL/SDK/串口/蓝牙接口的设备",
        "支持任何能通过 LSL 发布原始 EEG 数据的设备；当前已完成 BSense-R → BioMultiLite → LSL 链路验证，"
        "OpenBCI、Emotiv、NeuroSky 等品牌需按其 SDK 逐台适配与验收",
    ),
    (
        "只要设备能提供 SDK、串口或蓝牙接口即可接入，满足",
        "只要设备能通过 LSL 发布原始 EEG 数据即可接入（当前已验证 BSense-R 链路），满足",
    ),
    (
        "信号处理：MNE-Python、SciPy 与 NumPy 覆盖滤波、频谱与伪迹处理需求。",
        "信号处理：本项目以 NumPy 自研实现滤波、Welch 频谱与伪迹判定（不依赖 SciPy），MNE-Python 等工具作为离线复核手段。",
    ),
]

DETAIL_TABLE_24 = [
    ["维度", "厂商自带应用", "通用脑电研究平台", "本方案（凝思）"],
    ["指标可核对", "只给分数，算法不公开", "给原始数据与算法，需自行实现", "每个指标同时给原始频带功率、基线统计量与 z 值"],
    ["结论可复算", "不可复算", "可复算但流程由使用者自行拼接", "处理链参数随窗留痕，报告带口径版本号，可沿链路回算"],
    ["量表与行为联动", "通常缺失", "需自行设计", "SAS/SDS 与 SART/PVT-B 内置，三类证据联合判定"],
    ["不一致处理", "不适用", "由研究者自行判断", "明确标注结果不一致、建议复测，不强行合并分数"],
    ["边界说明", "常见夸大表述", "学术表述", "报告固定输出边界声明与非诊断提示"],
]

DETAIL_TABLE_25 = [
    ["需求来源", "主要发现", "在系统中的落点", "明确不承诺"],
    ["校园访谈（学生、辅导员）", "需要考前快速自查与短时放松，且必须能看懂指标含义", "约 10 分钟快速模式、指标说明与建议", "不做临床诊断、不做个体能力评价"],
    ["企业员工关怀访谈", "关注群体层面的状态分布，不接受对个人贴标签", "团体活动的汇总趋势与达标情况", "不提供个人排名与考核接口"],
    ["同类产品试用观察", "指标不透明、无法核对；缺少量表联动", "三类证据联合评估与证据回填", "不承诺与厂商分数可比"],
    ["既有采集工程试采记录", "伪迹与佩戴问题是主要数据损失来源", "逐窗质量门控与不可用原因记录", "不承诺在非受控环境下达到实验室质量"],
]

DETAIL_TABLE_73 = [
    ["界面页签", "用户可完成的事", "实现模块", "可验证产出"],
    ["状态指标", "查看专注度、放松度、认知负荷实时评分", "signal/indicators.py、app/ui.py", "每窗评分与相对基线 z 值写入报告"],
    ["实时波形与频谱", "核对原始波形与功率谱，确认信号正常", "signal/spectrum.py、signal/window.py", "去直流波形与 PSD 图，参数含 NFFT 与窗函数"],
    ["状态热力图", "按时间查看状态分档与低质量缺失窗", "monitoring/heatmap.py", "热力图 SVG 与缺失窗统计"],
    ["神经反馈训练", "按实时反馈完成分段训练并看到目标调整", "training/neurofeedback.py", "分段均值、达标时间占比、波动与下一段目标"],
    ["评估报告", "查看结论、建议与证据回填", "assessment/report.py", "报告 Markdown/JSON，含口径版本号"],
    ["历史趋势", "查看周/月趋势与可比性提示", "monitoring/history.py", "趋势 SVG 与不可比原因"],
]

DETAIL_TABLE_96 = [
    ["文档章节", "实现位置", "自动化用例"],
    ["6.3.1 工作流与四级处理链", "src/ningsi/signal/preprocess.py", "tests/test_preprocess_model.py::PreprocessTest"],
    ["6.4 模型训练与部署", "src/ningsi/models/logistic.py", "tests/test_preprocess_model.py::ModelTest"],
    ["6.6 状态热力图", "src/ningsi/monitoring/heatmap.py", "tests/test_assessment.py::HeatmapHistoryTest"],
    ["6.7 指标公式与量纲", "src/ningsi/signal/indicators.py、config.py", "tests/test_signal.py::IndicatorTest"],
    ["7.2.3 量表填写", "src/ningsi/scales/", "tests/test_scales_behavior.py::ScaleTest"],
    ["7.2.5 专注力行为任务", "src/ningsi/behavior/sart.py、pvt.py", "tests/test_scales_behavior.py::BehaviorTest"],
    ["7.2.10 评估报告", "src/ningsi/assessment/report.py", "tests/test_assessment.py::ReportTest"],
    ["7.2.11 历史与趋势", "src/ningsi/monitoring/history.py", "HeatmapHistoryTest::test_history_aggregation_and_comparability"],
    ["7.2.12 / 7.2.17 预警规则", "src/ningsi/monitoring/alerts.py", "tests/test_assessment.py::AlertTest"],
    ["8.6 自适应训练目标", "src/ningsi/training/neurofeedback.py", "tests/test_assessment.py::TrainingTest"],
    ["端到端会话", "src/ningsi/app/pipeline.py", "tests/test_pipeline.py::PipelineTest"],
]


CASE_STATUS = {
    "W-01": "部分实现（提示未发现流与安装指引，三类原因待补）",
    "W-02": "待开发（未做设备类型与通道数比对）",
    "W-07": "已实现（基线合格窗不足即拒绝出指标）",
    "W-08": "已实现，通过",
    "W-09": "已实现，通过",
    "W-10": "已实现，通过",
    "W-11": "已实现，通过",
    "W-12": "已实现，通过",
    "W-13": "已实现，通过",
    "W-14": "已实现，通过",
    "W-15": "已实现，通过",
    "W-16": "已实现，通过",
    "W-17": "已实现，通过",
    "W-18": "部分实现（缺答按 1 分计入并列出缺题号，未阻断提交）",
    "W-19": "已实现，通过",
    "W-20": "已实现，通过",
    "W-21": "已实现，通过",
    "W-22": "待开发（高分转介提示未实现）",
    "W-23": "已实现，通过",
    "W-24": "已实现，通过",
    "W-25": "部分实现（当前策略为下调目标，呼吸与坐姿提示未实现）",
    "W-26": "已实现，通过",
    "W-29": "待开发（尚未提供数据包导出）",
    "W-30": "待开发（模型加载未做版本登记校验）",
}

CAPABILITY_UPDATES = {
    "专注度与放松度指标": ("已在凝思中实现：β/θ、α/β 频带比值相对个体基线做 z 标准化后映射为 0–1 评分", "已实现（本轮新增模块）"),
    "心理量表集成": ("已在凝思中实现：SAS/SDS 各 20 题、反向题计分、粗分 × 1.25 与分级", "已实现（本轮新增模块）"),
    "神经反馈训练闭环": ("已在凝思中实现：分段统计、目标与保持时长自适应、训练前后同口径基线对比", "已实现（本轮新增模块）"),
    "历史记录与趋势分析": ("已在凝思中实现：JSONL 历史、周/月聚合与跨设备可比性保护", "已实现（本轮新增模块）"),
}


def update_case_table(doc) -> None:
    for table in doc.tables:
        header = [c.text.strip() for c in table.rows[0].cells]
        if header[:2] == ["编号", "模块"] and "结论" in header:
            index = header.index("结论")
            changed = 0
            for row in table.rows[1:]:
                code = row.cells[0].text.strip()
                if code in CASE_STATUS:
                    row.cells[index].text = CASE_STATUS[code]
                    for paragraph in row.cells[index].paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(9)
                    changed += 1
            log(f"  [更新] 白盒用例表结论列：{changed} 条按实现状态标注")
            return
    log("  [缺失] 未找到白盒用例表")


def update_capability_table(doc) -> None:
    for table in doc.tables:
        header = [c.text.strip() for c in table.rows[0].cells]
        if "复用状态" in header and "实现要点" in header:
            point = header.index("实现要点")
            state = header.index("复用状态")
            changed = 0
            for row in table.rows[1:]:
                name = row.cells[0].text.strip()
                if row.cells[point].text.strip() == "尚未实现" and name in CAPABILITY_UPDATES:
                    detail, status = CAPABILITY_UPDATES[name]
                    row.cells[point].text = detail
                    row.cells[state].text = status
                    for cell in (row.cells[point], row.cells[state]):
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                run.font.size = Pt(9)
                    changed += 1
            log(f"  [更新] 已有能力复用状态表：{changed} 条改为已实现")
            return
    log("  [缺失] 未找到已有能力对照表")


def update_label_table(doc) -> None:
    for table in doc.tables:
        header = [c.text.strip() for c in table.rows[0].cells]
        if "标签来源" in header and "当前状态" in header:
            state = header.index("当前状态")
            for row in table.rows[1:]:
                if row.cells[0].text.strip() == "专注力":
                    row.cells[state].text = "已实现（SART/PVT-B 试次标签 + logistic-v1 基线模型与被试级划分）"
                    for paragraph in row.cells[state].paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(9)
                    log("  [更新] 三类目标标签表：专注力一行改为已实现")
                    return
    log("  [缺失] 未找到三类目标标签表")


def update_material_table(doc) -> None:
    for table in doc.tables:
        header = [c.text.strip() for c in table.rows[0].cells]
        if header[:2] == ["序号", "材料"] and "完成状态" in header:
            index = header.index("完成状态")
            for row in table.rows[1:]:
                name = row.cells[1].text.strip()
                if "PPT" in name:
                    row.cells[index].text = "骨架已产出（deliverables/A09-凝思-项目简介.pptx，含逐页讲稿）"
                elif "演示视频" in name:
                    row.cells[index].text = "分镜脚本已完成，待录制（deliverables/演示视频分镜脚本.md）"
                elif "产品使用说明" in name:
                    row.cells[index].text = "已完成（deliverables/A09凝思_产品使用说明.docx，含系统架构与流程说明）"
                elif "概要介绍" in name:
                    row.cells[index].text = "已提供 v2 口径修订版（deliverables/）"
                elif "详细方案" in name:
                    row.cells[index].text = "已提供 v2 口径修订版（deliverables/）"
                for paragraph in row.cells[index].paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
            log("  [更新] 表 10.3 五项必交材料完成状态列")
            return
    log("  [缺失] 未找到必交材料清单表")


def revise_detail(source: Path) -> Path:
    log(f"\n=== 详细方案：{source.name} ===")
    doc = Document(str(source))
    prepare_cover(doc, "项目详细方案")
    replace_all(doc, DETAIL_REPLACEMENTS)

    heading_template = find(doc, exact="9 系统测试")
    body_template = find(doc, contains="测试策略分四层")
    caption_template = find(doc, startswith="表 9.1")

    # 2 章：补差异化对照与需求来源
    anchor = find(doc, exact="3 目标与服务模型")
    if anchor is not None and heading_template is not None and body_template is not None and caption_template is not None:
        block = [
            clone_paragraph(heading_template, "2.4.3 现有平台与本方案差异化对照"),
            clone_paragraph(body_template, "与厂商自带应用和通用脑电研究平台相比，本方案的差异不在指标数量，"
                                          "而在指标可核对、结论可复算、边界可说明三件事上，对照如表 2.4 所示。"),
            clone_paragraph(caption_template, "表 2.4 差异化定位对照表"),
            build_table(doc, DETAIL_TABLE_24),
            spacer(doc),
            clone_paragraph(heading_template, "2.5 需求来源与调研口径"),
            clone_paragraph(body_template, "本章的用户画像与竞品判断来自三条渠道：校园与企业场景访谈、同类消费级产品公开资料与配套应用的试用观察、"
                                          "以及本团队既有采集工程在真实设备上的试采记录。三条渠道只用于确定要做什么与不做什么，"
                                          "不参与任何个体结论。"),
            clone_paragraph(body_template, "调研结论按可执行性筛选后进入需求池，无法用现有硬件与样本量支撑的需求明确列为不承诺项，"
                                          "避免把看起来先进却无法交付的功能写进方案。需求来源与落点如表 2.5 所示。"),
            clone_paragraph(caption_template, "表 2.5 需求来源、调研结论与系统落点"),
            build_table(doc, DETAIL_TABLE_25),
            spacer(doc),
        ]
        insert_block(anchor, block, before=True)
        log("  [新增] 第 2 章：2.4.3 + 2.5 两节、2 张表")
    else:
        log("  [缺失] 第 2 章插入锚点未找到")

    # 6.4 / 6.5 口径补充
    anchor_64 = find(doc, startswith="6.4.1")
    if anchor_64 is not None and body_template is not None:
        insert_block(anchor_64, [clone_paragraph(
            body_template,
            "本版本已交付可训练、可部署的基线模型：输入为 6 维频带特征（θ/α/β 相对功率与三个比值），采用逻辑回归，"
            "按被试做 6:2:2 划分以避免窗口泄漏，输出准确率、灵敏度、特异度、AUC 与混淆矩阵，"
            "权重与口径以 JSON 版本化保存（logistic-v1）。深度模型（如 EEGNet）属下一阶段，"
            "上游数据集工程已具备窗口导出能力，可在不改变采集与质量流程的前提下替换模型。")], before=True)
        log("  [新增] 6.4 模型训练与部署：补充已交付基线模型的说明")
    anchor_65 = find(doc, startswith="6.5.1")
    if anchor_65 is not None and body_template is not None:
        insert_block(anchor_65, [clone_paragraph(
            body_template,
            "本版本对压力的输出是「压力相关维度指标偏高」的可操作提示，由认知负荷与放松度评分共同给出；"
            "情绪分类模型尚未交付，属下一阶段工作。本节给出的可操作定义与标签来源正是该阶段采集与验证的依据，"
            "因此当前报告只陈述观察到的指标变化，不为情绪状态命名。")], before=True)
        log("  [新增] 6.5 压力与情绪识别：标注当前边界")

    # 7 章：界面与实现对应
    anchor7 = find(doc, exact="8 项目亮点")
    if anchor7 is not None and heading_template is not None and body_template is not None and caption_template is not None:
        block = [
            clone_paragraph(heading_template, "7.3 界面与实现的对应关系"),
            clone_paragraph(body_template, "为避免出现文档写了界面而代码里找不到的情况，各页签与实现模块、可验证产出一一对应，"
                                          "如表 7.9 所示。表中每一行都能在代码仓库与自动化用例中定位。"),
            clone_paragraph(caption_template, "表 7.9 界面页签、实现模块与可验证产出对照表"),
            build_table(doc, DETAIL_TABLE_73),
            spacer(doc),
            clone_paragraph(heading_template, "7.4 一次完整会话的界面流转"),
            clone_paragraph(body_template, "一次完整会话在界面上的流转为：会话设置（匿名编号、Session、Run 与数据根目录）→ 设备接入与质检 → "
                                          "量表填写 → 睁眼与闭眼基线 → 行为任务 → 实时指标与热力图 → 联合评估报告 → "
                                          "神经反馈训练（含训练后基线）→ 历史趋势。全过程在同一窗口内完成，不需要在多个程序之间搬运数据。"),
            clone_paragraph(body_template, "训练之外的时间全部留有质量记录：质检与基线阶段记录不可用窗原因，任务阶段的低质量窗不计入指标与预警计时，"
                                          "报告在采集质量一节给出可用窗比例与排除原因分布，复测时可直接定位问题出现在哪个环节。"),
        ]
        insert_block(anchor7, block, before=True)
        log("  [新增] 第 7 章：7.3 + 7.4 两节、1 张表")
    else:
        log("  [缺失] 第 7 章插入锚点未找到")

    # 9 章：测试与证据索引
    anchor9 = find(doc, exact="10 项目总结")
    if anchor9 is not None and heading_template is not None and body_template is not None and caption_template is not None:
        block = [
            clone_paragraph(heading_template, "9.5 自动化测试与证据索引"),
            clone_paragraph(body_template, "除人工验收外，全部核心规则都有自动化用例覆盖，共 51 个用例，覆盖信号处理、频谱与频带、伪迹与质量、"
                                          "基线归一化、三指标、量表计分、行为任务、联合评估、报告证据回填、预警计时、热力图、历史趋势、"
                                          "训练自适应、模型训练与端到端会话。文档章节与实现位置、用例的对应关系如表 9.6 所示。"),
            clone_paragraph(caption_template, "表 9.6 文档章节、实现位置与自动化用例对照表"),
            build_table(doc, DETAIL_TABLE_96),
            spacer(doc),
            clone_paragraph(body_template, "复现方式：在项目根目录执行 scripts/verify.ps1（等价于逐条运行单元测试再加一次端到端演示），"
                                          "命令会打印用例统计、会话摘要与全部产物路径；模型训练结果（被试级划分、留出被试准确率与 AUC）"
                                          "随报告一并输出到 models/classifier.json，可直接作为验证证据提交。"),
        ]
        insert_block(anchor9, block, before=True)
        log("  [新增] 第 9 章：9.5 测试与证据索引、1 张表")
    else:
        log("  [缺失] 第 9 章插入锚点未找到")

    # 9 章与附录：把"尚未实现 / 待开发确认"更新为代码里的真实状态
    update_case_table(doc)
    update_capability_table(doc)
    update_label_table(doc)

    rewrites = [
        ("问题一：新模块尚未实现与测试",
         "问题一：新模块已完成开发，实测验证仍待完成。 专注度与放松度指标、量表集成、联合评估、神经反馈训练与趋势分析均已实现，"
         "并有 51 个自动化用例覆盖（见 9.5 节）；但尚未在真实设备与足够样本上完成效度验证。本方案不虚报验证结论，"
         "实测执行结果将在完成后随产品使用说明与测试记录提交。"),
        ("关于结论列的说明",
         "关于结论列的说明：标注“已实现，通过”表示该用例在本轮交付代码中已实现并有自动化用例覆盖（见 9.5 节）；"
         "标注“部分实现”表示主路径可用但边界行为与表中描述仍有差异，差异写在该行结论中；标注“待开发”表示仍属下一阶段。"
         "本方案不预设未执行测试的结论，执行记录将随产品使用说明文档提交。"),
        ("赛题要求的提交物为五项必交材料",
         "赛题要求的提交物为五项必交材料，即项目概要介绍、项目简介 PPT、项目详细方案、项目演示视频、产品使用说明文档；"
         "其他补充材料按赛题属自愿提交项，不作为必交。"),
        ("清单中已提供的两项材料为本轮可核验版本",
         "截至本轮修订：项目概要介绍与项目详细方案已提供 v2 口径修订版（封面、口径与篇幅配比同步更新）；"
         "项目简介 PPT 骨架已产出并附逐页讲稿，待补充界面截图；产品使用说明文档已完成，含系统架构与流程说明；"
         "项目演示视频分镜脚本已完成，待按脚本录制。完成状态以表 10.3 为准，演示视频时长 5–8 分钟为团队建议，赛题未规定。"),
    ]
    for prefix, new_text in rewrites:
        target = find(doc, startswith=prefix)
        if target is not None:
            set_paragraph_text(target, new_text)
            log(f"  [更新] {prefix[:16]}…")
        else:
            log(f"  [缺失] {prefix[:16]}…")

    update_material_table(doc)

    set_update_fields(doc)
    return save(doc, "A09-凝思-项目详细方案-v2.docx")


def main() -> int:
    import sys
    global SUFFIX
    if len(sys.argv) > 1 and sys.argv[1].startswith("--suffix"):
        SUFFIX = sys.argv[1].split("=", 1)[-1] if "=" in sys.argv[1] else sys.argv[2]
        log(f"输出后缀：{SUFFIX}")
    overview = SOURCE / "A09-凝思-项目概要介绍-最终版.docx"
    detail = SOURCE / "A09-凝思-项目详细方案-最终版.docx"
    for path in (overview, detail):
        if not path.exists():
            log(f"找不到源文件：{path}")
            return 1
    revise_overview(overview)
    revise_detail(detail)
    log("\n完成。注意：Word 打开新文件时会提示更新域，选中目录按 F9 也可刷新页码。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
