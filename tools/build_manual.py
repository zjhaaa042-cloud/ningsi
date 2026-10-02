"""把 docs/PRODUCT_GUIDE.md 转成可直接提交的 Word 版《产品使用说明文档》。

对应赛题提交材料「（5）① 产品使用说明文档（系统架构与流程说明）」，
输出文件名与详细方案表 10.3 的建议文件名一致：A09凝思_产品使用说明.docx。
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]          # D:/Desktop/ninsi
SOURCE = ROOT / "ningsi" / "docs" / "PRODUCT_GUIDE.md"
OUT = ROOT / "deliverables"
OUT.mkdir(parents=True, exist_ok=True)

BODY_FONT = "微软雅黑"
CODE_FONT = "Consolas"
ACCENT = RGBColor(0x2B, 0x3A, 0x67)
GREY = RGBColor(0x44, 0x4C, 0x57)
CODE_FILL = "F2F4F7"
HEADER_FILL = "E8EDF4"


def set_run_font(run, name=BODY_FONT, size=None, bold=None, color=None, italic=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), name)
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction_node = OxmlElement("w:instrText")
    instruction_node.set(qn("xml:space"), "preserve")
    instruction_node.text = instruction
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instruction_node)
    run._r.append(end)


def shade_paragraph(paragraph, fill: str) -> None:
    element = OxmlElement("w:shd")
    element.set(qn("w:val"), "clear")
    element.set(qn("w:fill"), fill)
    paragraph._p.get_or_add_pPr().append(element)


def shade_cell(cell, fill: str) -> None:
    element = OxmlElement("w:shd")
    element.set(qn("w:val"), "clear")
    element.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(element)


def style_document(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.28
    for name, size in (("Heading 1", 17), ("Heading 2", 14), ("Heading 3", 12)):
        style = doc.styles[name]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = ACCENT
        style.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 10)
        style.paragraph_format.space_after = Pt(6)
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.4)
    section.top_margin = section.bottom_margin = Cm(2.2)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    prefix = footer.add_run("第 ")
    settle(prefix, 9)
    add_field(footer, "PAGE")
    middle = footer.add_run(" 页  共 ")
    settle(middle, 9)
    add_field(footer, "NUMPAGES")
    tail = footer.add_run(" 页")
    settle(tail, 9)


def settle(run, size):
    set_run_font(run, BODY_FONT, size, color=GREY)


def add_cover(doc: Document) -> None:
    for _ in range(4):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(title.add_run("产品使用说明文档"), BODY_FONT, 30, bold=True, color=ACCENT)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(subtitle.add_run("系统架构与流程说明"), BODY_FONT, 16, color=GREY)
    doc.add_paragraph()
    slogan = doc.add_paragraph()
    slogan.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(slogan.add_run("凝思（Ningsi）—— 每一段专注，都看得见"), BODY_FONT, 14, color=ACCENT)
    for _ in range(6):
        doc.add_paragraph()
    for text in ("赛题编号：A09", "命题企业：杭州金扬智能科技有限公司",
                 "文档版本：V1.0（与代码版本 ningsi 0.1.0 一致）", "团队名称：胆double天队",
                 "汇报人：____________（提交前填写）", "日期：2026 年 9 月"):
        line = doc.add_paragraph()
        line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(line.add_run(text), BODY_FONT, 12)
    doc.paragraphs[-1].runs[-1].add_break(WD_BREAK.PAGE)


def add_toc(doc: Document) -> None:
    heading = doc.add_paragraph()
    set_run_font(heading.add_run("目　　录"), BODY_FONT, 16, bold=True, color=ACCENT)
    paragraph = doc.add_paragraph()
    add_field(paragraph, 'TOC \\o "1-3" \\h \\z \\u')
    hint = doc.add_paragraph()
    set_run_font(hint.add_run("（打开文档时如提示更新域，请选择“是”；也可选中目录按 F9 刷新页码）"),
                 BODY_FONT, 9, color=GREY)
    hint.runs[-1].add_break(WD_BREAK.PAGE)


INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def add_inline(paragraph, text: str, size: float = 10.5) -> None:
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**"):
            set_run_font(paragraph.add_run(piece[2:-2]), BODY_FONT, size, bold=True)
        elif piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            set_run_font(run, CODE_FONT, size - 1.5)
        else:
            set_run_font(paragraph.add_run(piece), BODY_FONT, size)


def add_table(doc: Document, rows) -> None:
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for column_index, value in enumerate(row):
            cell = table.cell(row_index, column_index)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            add_inline(paragraph, value.strip(), 9.5)
            if row_index == 0:
                shade_cell(cell, HEADER_FILL)
                for run in paragraph.runs:
                    run.font.bold = True
    doc.add_paragraph()


def convert(markdown: str, doc: Document) -> dict:
    stats = {"headings": 0, "tables": 0, "codes": 0, "lists": 0, "paragraphs": 0}
    lines = markdown.split("\n")
    index = 0
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if stripped.startswith("```"):
            index += 1
            block = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(8)
            shade_paragraph(paragraph, CODE_FILL)
            for position, code_line in enumerate(block):
                run = paragraph.add_run(code_line)
                set_run_font(run, CODE_FONT, 9)
                if position != len(block) - 1:
                    run.add_break()
            stats["codes"] += 1
        elif stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            if level == 1:
                stats["headings"] += 1
            else:
                heading = doc.add_paragraph(style=f"Heading {min(3, level - 1)}")
                add_inline(heading, text, {2: 14, 3: 12}.get(min(3, level - 1), 12))
                for run in heading.runs:
                    run.font.bold = True
                    run.font.color.rgb = ACCENT
                stats["headings"] += 1
        elif stripped.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                if not all(set(cell) <= set("-: ") for cell in cells):
                    rows.append(cells)
                index += 1
            index -= 1
            if rows:
                add_table(doc, rows)
                stats["tables"] += 1
        elif stripped.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            add_inline(paragraph, stripped[2:])
            stats["lists"] += 1
        elif re.match(r"^\d+\. ", stripped):
            paragraph = doc.add_paragraph(style="List Number")
            add_inline(paragraph, re.sub(r"^\d+\. ", "", stripped))
            stats["lists"] += 1
        elif stripped.startswith(">"):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.6)
            shade_paragraph(paragraph, CODE_FILL)
            add_inline(paragraph, stripped.lstrip("> ").strip(), 10)
            stats["paragraphs"] += 1
        elif stripped in ("---", "***", "___"):
            pass
        elif stripped:
            paragraph = doc.add_paragraph()
            add_inline(paragraph, stripped)
            stats["paragraphs"] += 1
        index += 1
    return stats


def enable_update_fields(doc: Document) -> None:
    element = OxmlElement("w:updateFields")
    element.set(qn("w:val"), "true")
    doc.settings.element.append(element)


def main() -> int:
    markdown = SOURCE.read_text(encoding="utf-8")
    document = Document()
    style_document(document)
    add_cover(document)
    add_toc(document)
    stats = convert(markdown, document)
    enable_update_fields(document)
    target = OUT / "A09凝思_产品使用说明.docx"
    document.save(target)
    print(f"  已生成 {target}")
    print(f"  统计：标题 {stats['headings']} 个，表格 {stats['tables']} 张，代码块 {stats['codes']} 个，"
          f"列表 {stats['lists']} 条，段落 {stats['paragraphs']} 段")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
